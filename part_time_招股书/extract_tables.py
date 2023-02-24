import json
import re
import jieba
import pandas as pd
from docx import Document
import docx
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table, _Row
from docx.text.paragraph import Paragraph
from tqdm import tqdm
import os
import itertools
import time
from multiprocessing import Pool
import logging
from pprint import pprint
from openpyxl import Workbook
from collections import defaultdict
from copy import  deepcopy
from typing import Sequence

def merge_files(all_files):
    all_files = [('-'.join(i.split('-')[:-3]), i) for i in all_files]
    all_files = sorted(all_files, key=lambda x:x[0])
    all_files = [(k, list(g)) for k,g in itertools.groupby(all_files, key=lambda x:x[0])]
    all_files = [(k, [i[1] for i in g]) for k,g in all_files]
    return all_files

def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    elif isinstance(parent, _Row):
        parent_elm = parent._tr
    else:
        raise ValueError("something's not right")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def recognize_table(table: Table):
    result = []
    for row in table.rows:
        cells = row.cells
        result.append([re.sub(r'\s+', '', i.text) for i in cells])
    return result

def read_merged_doc(batch_files):
    final_result = []
    for doc_file in batch_files:
        doc = Document(doc_file)
        cur_doc_content = []
        for block in iter_block_items(doc):
            # read Paragraph
            if isinstance(block, Paragraph) and block.text:
                if block.text == '四、发行人报告期的主要财务数据和财务指标':
                    a = 1
                cur_doc_content.append(block.text)
            # read table
            elif isinstance(block, Table):
                table_content = recognize_table(block)
                cur_doc_content.append(('table', table_content))
        cur_doc_content = [i for i in cur_doc_content if i]
        final_result.extend(cur_doc_content)
    return final_result

################################################# 表格抽取 ##########################################################
'''
主营业务收入
'''
def extract_table_given_key_zhuyingshourugoucheng(filename, total_doc_content):
    '''
    主营业务构成
    :param filename:
    :param total_doc_content:
    :return:
    '''
    key = '主营业务构成'
    patterns = [
        re.compile(r'营业收入[的]?构成'),
        re.compile(r'营业收入分析'),
    ]

    def is_table_zhuyingyewu(table):
        first_col = [i[0] for i in table]
        # cells = [i for j in table for i in j]
        if first_col[:3] == ['项目', '项目', '主营业务收入'] \
                or first_col[:3] == ['项目', '主营业务收入', '其他业务收入']\
                or first_col[:3] == ['项目', '主营业务收入', '其它业务收入']\
                or first_col[:3] == ['收入类型', '主营业务收入', '其它业务收入']\
                or first_col[:3] == ['项目', '项目', '主营业务']\
                or first_col[-2:] == ['其它业务收入','合计']\
                or first_col[-2:] == ['其它业务','合计']\
                or first_col[-2:] == ['其他业务','合计']\
                or first_col[2:4] == ['主营业务收入','其它业务收入']\
                or first_col[2:4] == ['主营业务收入','其他业务收入']\
                :
            return True
        return False

    find = False
    find_key_lines = []
    for lineno, line in enumerate(total_doc_content):
        if isinstance(line, str):
            matched_pat = [pat.search(line.replace(' ', '')) for pat in patterns]
            if any(matched_pat):
                find = True
                # print(f'{filename} {lineno} {line}')
                find_key_lines.append(lineno)

    if len(find_key_lines) == 0:
        print(filename, key, '没找到')
        find_key_lines_new = []
    elif len(find_key_lines) == 1:
        find_key_lines_new = find_key_lines
    else:
        find_key_lines_new = [find_key_lines[0]]
        for idx, lineno in enumerate(find_key_lines[1:], start=1):
            if lineno - find_key_lines[idx-1] == 1:
                # find_key_lines_new[-1] = lineno
                continue
            else:
                find_key_lines_new.append(lineno)
    result = []
    if find_key_lines_new:
        for lineno in find_key_lines_new:
            result.append('#'*100)
            result.extend(total_doc_content[lineno:lineno+6])

    return result

'''
主营业务按产品、地域等划分
'''
def to_string(i):
    if isinstance(i, str):
        return i
    elif isinstance(i, tuple):
        table = i[1]
        table_str = [' | '.join(row) for row in table]
        table_str = '\n'.join(table_str)
        return table_str
def extract_table_given_key_zhuying_different_part(filename, total_doc_content):
    ''''''
    all_result = []
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple):continue
        if re.search(r'[.)）、]主营业务.{0,7}(产品|地域|区域|地区|季节|结构|构成|业务类别).{0,8}$', re.sub(r'\s', '', line)) or re.search(r'[.)）、].{0,4}(产品|地区|区域|季节|地域).{0,4}主营业务.{0,5}$', re.sub(r'\s', '', line)):
            lines_below = total_doc_content[idx:idx+8]
            all_result.append('#' * 100)
            all_result.extend(lines_below)

    # with open(f'tmp6_主营业务划分/{filename}.txt', mode='w', encoding='utf8') as f:
    #     f.write('\n'.join(all_result))
    return all_result

'''
主要财务数据
'''
def extract_table_given_key_zhuyaocaiwushuju(filename, total_doc_content):
    '''
    主要财务收入
    :param filename:
    :param total_doc_cotent:
    :return:
    '''
    key = '主要财务数据和财务指标'
    patterns = [
        re.compile(r'^三(.+?)主要财务数据[和及与].{0,2}?财务指标'),
        re.compile(r'^三(.+?)主要财务数据$'),
        re.compile(r'^.{0,15}主要财务数据.{0,5}财务指标.{0,5}$')
        # ('营业收入构成', re.compile(r'营业收入[的]?构成')),
        # ('营业收入构成', re.compile(r'营业收入分析')),
    ]

    lineno_key = []
    result = []
    for lineno, line in enumerate(total_doc_content):
        if isinstance(line, str):
            matched_pat = [pat.search(line.replace(' ', '')) for pat in patterns]
            if any(matched_pat):
                lineno_key.append(lineno)
                # break
    if not lineno_key:
        print(filename, key, '没找到')
    else:
        for lineno in lineno_key:
            result.append('#'*100)
            result.extend(total_doc_content[lineno:lineno+6])
    return result

'''
本次发行概况
'''
def extract_table_given_key_bencifaxinggaikuang(filename, total_doc_content):
    key = '本次发行概况'
    patterns = [
        re.compile(r'第三节\s+本次发行概况')
    ]
    lineno_key = []
    result = []
    for lineno, line in enumerate(total_doc_content):
        if isinstance(line, str):
            matched_pat = [pat.search(line) for pat in patterns]
            if any(matched_pat):
                lineno_key.append(lineno)
                # break
    if not lineno_key:
        print(filename, key, '没找到')
    else:
        for lineno in lineno_key:
            result.append('#'*100)
            result.extend(total_doc_content[lineno:lineno+5])
    return result

'''
发行人及本次发行的中介机构基本情况
'''
def extract_table_given_key_faxingrenhefaxingzhongjie(filename, total_doc_content):
    key = '发行人及本次发行的中介机构基本情况'
    patterns = [
        re.compile(r'一、发行人及本次发行的中介机构基本情况'),
        re.compile(r'一、本次发行的有关当事人基本情况'),
        re.compile(r'一、发行人及本次发行相关中介机构基本情况')
    ]
    end = re.compile(r'二、本次发行概况|二、本次发行情况')
    lineno_key = None
    for lineno, line in enumerate(total_doc_content):
        if isinstance(line, str):
            matched_pat = [pat.search(line.replace(' ', '')) for pat in patterns]
            if any(matched_pat):
                lineno_key = lineno
                break
    if not lineno_key:
        print(filename, key, '没找到')
    # else:
    #     end_line = None
    #     for lineno_end in range(lineno_key, lineno_key+10):
    #         line = total_doc_content[lineno_end]
    #         if isinstance(line, str) and end.search(line.replace(' ', '')):
    #             end_line = lineno_end
    #             break
    #     if end_line:
    #         result = []
    #         for lineno in range(lineno_key, end_line):
    #             if isinstance(total_doc_content[lineno], tuple):
    #                 result.extend(total_doc_content[lineno][1])
    #         if result:
    #             return result
    # print(filename, key, 'end 没找到')
    return None

def match_in_list(pat, lst):
    matched = [pat.search(i) for i in lst]
    return any(matched)

def extract_table_given_key_faxingrenhefaxingzhongjieV2(filename, total_doc_content):
    key = '发行人及本次发行的中介机构基本情况'
    fileds_person = [
        ('发行人名称', re.compile(r'发行人名称|中文名称')),
        ('成立日期', re.compile('股份公司成立日期|有限公司成立日期|成立日期')),
        ('注册资本', re.compile('注册资本')),
        ('法定代表人', re.compile(r'法定代表人')),
        ('注册地址', re.compile(r'注册地址')),
        ('主要生产经营地址', re.compile('主要生产经营地址')),
        ('控股股东', re.compile(r'控股股东')),
        ('实际控制人', re.compile(r'实际控制人')),
        ('行业分类', re.compile(r'行业分类')),
        ('在其他交易所(申请)挂牌或上市的情况', re.compile(r'在其他交易所(申请)挂牌或上市的情况|在其他交易场所(申请)挂牌或上市的情况')),
       ]
    fileds_jigou = [('保荐人',  re.compile(r'保荐人|联合保荐机构')),
        ('主承销商', re.compile(r'主承销商')),
        ('发行人律师', re.compile(r'发行人律师')),
        ('其他承销机构', re.compile(r'其他承销机构')),
        ('审计机构', re.compile(r'审计机构暨验资复核机构|审计机构')),
        ('评估机构', re.compile(r'评估机构(二)|评估机构(如有)|评估机构')),

    ]
    tables = [i[1] for i in total_doc_content if isinstance(i, tuple)]
    matched_table_person = []
    matched_table_jigou = []
    for table in tables:
        all_items = {i for j in table for i in j}
        contain_fileds_person = [i for i in fileds_person if match_in_list(i[1], all_items)]
        if len(contain_fileds_person) / len(fileds_person) > 0.5:
            matched_table_person.append(table)
        contain_fileds_jigou = [i for i in fileds_jigou if match_in_list(i[1], all_items)]
        if len(contain_fileds_jigou) / len(fileds_jigou) > 0.5:
            matched_table_jigou.append(table)
    if len(matched_table_person) > 2:
        print(filename, 'person > 2')
    if len(matched_table_person) == 0:
        print(filename, 'person == 0')
    if len(matched_table_jigou) > 2:
        print(filename, 'jigou > 2')
    if len(matched_table_jigou) == 0:
        print(filename, 'jigou == 0')

    result = {}

    for table in matched_table_person:
        for line in table:
            for col_no, col in enumerate(line):
                for k, pat in fileds_person:
                    if pat.search(col):
                        try:
                            key = pat.search(col).group(0)
                            val = line[col_no+1]
                            result[key] = val
                        except Exception as e:
                            print('error: ', key, line)

    for table in matched_table_jigou:
        for line in table:
            for col_no, col in enumerate(line):
                for k, pat in fileds_jigou:
                    if pat.search(col):
                        try:
                            key = pat.search(col).group(0)
                            val = line[col_no + 1]
                            result[key] = val
                        except Exception as e:
                            print('error: ', key, line)
    result = ['#'*100, ('table', [[k,v] for k,v in result.items()])]
    return result

'''
财务会计信息抽取
'''
def post_process2(content):
    pat = re.compile(
        r'''
        ^.{0,7}((母公司|合并).{0,3})?负债表.{0,4}$|
        ^.{0,7}((母公司|合并).{0,3})?利润表.{0,4}$|
        ^.{0,7}((母公司|合并).{0,3})?现金流.{0,4}$
        ''', re.VERBOSE
    )
    keys = [line if (isinstance(line, str) and pat.search(line) and '续' not in line) else None for line in content]
    for idx, _line in enumerate(content):
        line = list(_line) if isinstance(_line, tuple) else deepcopy(_line)
        if isinstance(line, list):
            tb = line[1]
            first_line = '|'.join(tb[0])
            for i in range(1, 6):
                pre_idx = idx - i
                if isinstance(content[pre_idx], str) and keys[pre_idx] is not None:
                    tb = [[f'table_{keys[pre_idx]}']+row for row in tb]
                    line[1] = tb
                    break
                elif isinstance(content[pre_idx], list):
                    if i <= 3:
                        head = deepcopy(content[pre_idx][1][0])
                        if head[0].startswith('table') or head[0] == 'None':
                            head_0 = head.pop(0)
                        else:
                            head_0 = 'None'
                        tb = [[head_0] + row for row in tb]
                        line[1] = tb
                    else:
                        tb = [['None']+row for row in tb]
                        line[1] = tb
                    break
            content[idx] = line
    return content

def extract_table_given_key_caiwukuaiji_xinxi(filename, total_doc_content):
    key = '财务会计信息与管理层分析'
    chapter_pat = re.compile(r'^第八节\s*财务会计信息与管理层.{0,4}分析$')
    chapter_match = [idx for idx, i in enumerate(total_doc_content)  if isinstance(i, str) and chapter_pat.search(i.strip())]
    if not chapter_match:
        print(f'not found {filename}')
        return []
    elif len(chapter_match) > 1:
        print(f'too many {filename} {len(chapter_match)}')
        return []
    else:
        idx = chapter_match[0]
        lines_below = total_doc_content[idx:idx+100]
        result = post_process2(lines_below)

        return result

'''
发行前后公司股本情况
'''
def is_guben_bianhua(table):
    # key_words = re.compile(r'发行前|发行后|股份数|股数|比例|占比|股东|股份')
    key_words_and_nums = [
        (re.compile(r'发行前'), 1),
        (re.compile(r'发行后'), 1),
        (re.compile(r'股份数|股数|股份|股东'), 2),
        (re.compile(r'占比|比例|股数[(（]%[)）]'), 2)
    ]
    head2_rows = [i for j in table[:2] for i in j]
    pattern_matched_num = 0
    for pat, num in key_words_and_nums:
        matched = [1 for i in head2_rows if pat.search(i)]
        if sum(matched) >= num:
            pattern_matched_num += 1
    if pattern_matched_num == len(key_words_and_nums):
        return True
    else:
        return False

def get_float(num):
    num = re.sub(r'%|\s|,', '', num)
    try:
        return float(num)
    except:
        return None

def line2dict(line):
    name, num_before, rate_before, num_after, rate_after = line
    num_before = get_float(num_before)
    num_after = get_float(num_after)
    rate_before = get_float(rate_before)
    rate_after = get_float(rate_after)
    if num_before is not None and num_after is not None and rate_after is not None and rate_before is not None:
        return {'持股人':name, '发行前持股数':num_before, '发行前持股比例':rate_before, '发行后持股数':num_after, '发行后持股比例':rate_after}
    else:
        return {}

def extract_table_given_key_guben_bianhua(filename, total_doc_content):
    head_key = re.compile(r'本次发行前后公司股本情况')
    # tables = [i[1] for i in total_doc_content if isinstance(i, tuple)]
    guben_bianhua_table = []
    result = []
    for idx, line in enumerate(total_doc_content):
        # if line == '八、发行人股本情况':
        #     print()
        if isinstance(line, tuple) and is_guben_bianhua(line[1]):
            # guben_bianhua_table.extend(line[1])
            guben_bianhua_table.append(idx)
    if not guben_bianhua_table:
        print('股本变化没找到：',filename)
    else:
        # with open(f'tmp3_股本变化分析/{filename}.txt', mode='w', encoding='utf8') as f:
        #     f.write('\n'.join([' | '.join(i) for i in guben_bianhua_table]))
        for lineno in guben_bianhua_table:
            result.append('#'*100)
            result.extend(total_doc_content[lineno-4:lineno+2])
    return result

'''
员工专业结构、学历结构
'''
def extract_table_given_key_yuangong_zhuanye_xueli_old(filename, total_doc_content):
    zhuanye_result = []
    xueli_result = []
    zhuanye_table = set()
    xueli_table = set()
    lineno_key = []
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple):continue
        if re.search(r'专业结构|专业构成', line):
            for i in range(idx+1, idx+3):
                if isinstance(total_doc_content[i], tuple):
                    zhuanye_table.add(i)
            a = 1
        elif re.search(r'学历结构|学历构成|受教育程度', line):
            for i in range(idx + 1, idx + 3):
                if isinstance(total_doc_content[i], tuple):
                    xueli_table.add(i)
    zhuanye_table = [total_doc_content[i][1] for i in zhuanye_table]
    # zhuanye_table = ['\n'.join([' | '.join(i) for i in table]) for table in zhuanye_table]
    # with open(f'tmp4_员工比例/{filename}_专业.txt', mode='w', encoding='utf8') as f:
    #     f.write('\n####################### 专业比例 ###########################\n'.join(zhuanye_table))
    xueli_table = [total_doc_content[i][1] for i in xueli_table]
    # xueli_table = ['\n'.join([' | '.join(i) for i in table]) for table in xueli_table]
    # with open(f'tmp4_员工比例/{filename}_学历.txt', mode='w', encoding='utf8') as f:
    #     f.write('\n####################### 学历比例 ###########################\n'.join(xueli_table))
    for table in zhuanye_table:
        zhuanye_result.extend(table)
    for table in xueli_table:
        xueli_result.extend(table)
    return {'人员学历构成':xueli_result, '人员专业构成':zhuanye_result}

def extract_table_given_key_yuangong_zhuanye_xueli(filename, total_doc_content):
    lineno_key = []
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple):continue
        if re.search(r'专业结构|专业构成', line):
            lineno_key.append(idx)
        elif re.search(r'学历结构|学历构成|受教育程度', line):
            lineno_key.append(idx)
    result = []
    for lineno in lineno_key:
        result.append('#'*100)
        result.extend(total_doc_content[lineno:lineno+4])
    return result

'''
前五大客户销售情况
'''
def extract_table_given_key_top_five_kehu(filename, total_doc_content):
    result = []
    pat = re.compile(r'报告期(内|各期)?.{0,7}五[大名]客户.{0,3}(销售.{0,3}情况|应收账款)'
                     r'|主要客户.{0,3}销售.{0,2}情况'
                     r'|应收账款.{0,3}五大客户'
                     r'|报告期(内|各期)?.{0,8}销售.{0,5}前五名客户'
                     r'|公司向前五[名大]客户.{0,4}销售情况'
                     r'|报告期.{0,8}前五[名大]客户.{0,3}情况$'
                     r'|^.{0,8}前五[大名]客户.{0,8}$', re.VERBOSE)
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple):continue
        line = re.sub(r'[(（].+?[)）]', '', line)
        if pat.search(line):
        # if re.search(r'报告期内前五大客户销售情况', line):
            below = total_doc_content[idx:idx+15]
            # below_tables = [i[1] for i in below if isinstance(i, tuple)][:3]
            # if not below_tables:continue
            # result.extend(below_tables)
            result.extend(below)
            # break
    return result

'''
前五大供应商
'''
def extract_table_given_key_top_five_gongyings(filename, total_doc_content):
    result = []
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple):continue
        line = re.sub(r'[(（].+?[)）]', '', line)
        key = [
            '报告期内前五大原材料供应商采购情况',
            '报告期内主要供应商情况',
            '报告期内，公司前五大外协供应商的采购金额和基本情况如下',
            '报告期内，公司向前五大供应商的采购情况',
            '公司报告期内前五名供应商的采购情况',
            '报告期各期末，公司预付账款前五名供应商如下',
            '报告期公司前五名供应商情况',
            '报告期内向前五名供应商采购情况',
            '报告期内，公司向前五大供应商的采购情况',
            '报告期内，公司前五大供应商情况如下',
            '报告期内，公司向前五名供应商的采购情况',
            '^.{0,3}前五名供应商采购情况$',
            '报告期内前五名供应商采购情况',
            '报告期内，公司前五名供应商的采购金额',
            '报告期内前五名供应商情况',
            '报告期内，公司向前五名供应商的采购情况如下',
            '报告期内，公司前五名原材料供应商采购情况如下',# ?
            '报告期内，公司向前五名供应商采购情况如下',
            '报告期，发行人主要原材料前五大供应商情况',
            '报告期内，公司对前五大供应商采购金额',
            '报告期各期，公司前五名材料供应商采购情况如下',
            '报告期内前五大供应商的具体情况如下：'
        ]
        if re.search(r'^.{0,3}报告期(内|各期)?.{0,10}(前五|主要).{0,3}(原材料)?供应商.{0,4}(采购)|^.{0,3}前五名供应商采购情况.{0,4}$|^.{0,4}报告期.{0,13}前五.{0,7}供应商.{0,3}情况.{0,9}$', line):
            # 预付款有问题
            below = total_doc_content[idx:idx+15]
            # below_tables = [i[1] for i in below if isinstance(i, tuple)][:3]
            # if not below_tables:continue
            # result.extend(below_tables)
            result.extend(below)
            # break
    return result

'''
毛利及毛利率分析
'''
def extract_table_given_key_maolilv(filename, total_doc_content):
    result = []
    next_line = 0
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple) or idx < next_line:
            continue
        if re.search(r'^.{0,4}、.{0,8}毛利.{0,3}毛利率分析.{0,8}$|^.{0,3}[(（][一二三四五六七八九十][)）].{0,8}毛利(率)?.{0,6}分析.{0,8}$|^.{0,5}(主营业务)?毛利(率)?.{0,6}(分析|情况|变动).{0,3}$|^.{0,3}、.{0,10}毛利率分析$|^.{0,4}毛利构成.{0,7}分析|^.{0,3}[(（][一二三四五六七八九十][)）].{0,12}毛利(率)?.{0,3}分析.{0,3}$', line):
            below = total_doc_content[idx:idx+35]
            result.extend(below)
            next_line = idx+35
    if not result:
        print(filename, "没找到毛利及毛利率分析")
    return result

'''
前10股东变化
'''
def extract_table_given_key_top10_gudong(filename, total_doc_content):
    result = []
    next_line = 0
    pat = re.compile(
        r'''
        ^.{0,15}十大.{0,5}股东.{0,15}$|
        ^.{0,15}前(十|\s*10\s*)名.{0,5}股东.{0,15}$|
        ^.{0,15}内资股股东.{0,10}$
        ''', re.VERBOSE)
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple) or idx < next_line:
            continue
        if pat.search(line):
            below = total_doc_content[idx:idx + 35]
            result.extend(below)
            next_line = idx + 35
    if not result:
        print(filename, "没找到前10股东变化")
    return result

'''
发行人符合科创属性、行业领域要求
'''
def extract_table_given_key_meet_demand(filename, total_doc_content):
    result = []
    pat = re.compile(
        r'''
        ^.{0,10}符合行业领域.{0,8}要求.{0,5}$|
        ^.{0,10}符合科创属性.{0,8}要求.{0,5}$|
        ^.{0,10}符合科创板行业定位.{0,8}$
        ''', re.VERBOSE)
    next_line = 0
    for idx, line in enumerate(total_doc_content):
        if isinstance(line, tuple) or idx < next_line:
            continue
        if pat.search(line):
            below = total_doc_content[idx:idx + 35]
            result.append('#'*50)
            result.extend(below)
            next_line = idx + 35
    if not result:
        print(filename, "没有找到 符合科创属性、行业领域要求")
    return result

'''
营业收入分析
'''
def  extract_table_given_key_yingyeshouwu_analize(filename, total_doc_content):
    result = []
    pat = re.compile(
        r'''
        ^.{0,5}经营成果分析.{0,5}$|
        ^.{0,5}营业收入分析.{0,5}$|
        ^.{0,5}营业收入及成本分部信息.{0,5}$
        ''',re.VERBOSE
    )
    matched_lines = [(lineno, i) for lineno, i in enumerate(total_doc_content) if isinstance(i, str) and pat.search(i)]
    if not matched_lines:
        print('没有找到 经营成果分析|营业收入分析')
        return result
    end = [(lineno, i) for lineno, i in enumerate(total_doc_content) if isinstance(i, str) and re.search(r'^.{0,5}资产质量分析.{0,8}$', i)]
    if end:
        end_lineno = end[0][0]
    else:
        end_lineno = matched_lines[0][0] + 100
    result.extend(total_doc_content[matched_lines[0][0]:end_lineno])
    return result

##################################################### 指标抽取 #############################################################
def get_zhibiao_caiwu(result_per_tb):
    for line in result_per_tb:
        ...

def extract_zhibiao(result):
    财务会计信息与管理层分析章节抽取 = get_zhibiao_caiwu(result['财务会计信息与管理层分析章节抽取'])

############################################ 表格抽取结果写文件 ############################################
def write_to_excel(result, filename, sub_dir):
    book = Workbook()
    for k, v in result.items():
        sheet = book.create_sheet(k)
        for line in v:
            if isinstance(line, str):
                sheet.append((line, ))
            elif isinstance(line, Sequence):
                for row in line[1]:
                    sheet.append(row)
    save_path = f'最新文件抽取结果/{sub_dir}/{filename}.xlsx'
    os.makedirs(os.path.dirname(save_path), exist_ok=True)
    book.save(save_path)

############################################## 业务与技术章节提取 ########################################################

def get_hierarchy_of_chapter(yewu_and_jishu):
    patterns = [
        ('0', re.compile(r'^第.{0,2}[节章]\s*业务.{0,2}技术')),
        ('1', re.compile(r'[一二三四五六七八九十]{1,2}.{0,1}、.{0,35}')),
        ('2', re.compile(r'[(（][一二三四五六七八九十]{1,2}[)）]')),
        ('3', re.compile(r'[0-9]{0,2}、.{0,30}'))
    ]
    levels = []
    cur_ = ['' for _ in patterns]
    for line in yewu_and_jishu:
        if not isinstance(line, str):
            levels.append([i for i in cur_])
        else:
            for idx in range(len(patterns)):
                if patterns[idx][1].match(line.replace(' ', '')):
                    cur_[idx] = line
                    for i in range(idx+1, len(patterns)):
                        cur_[i] = ''
            levels.append([i for i in cur_])
    yewu_and_jishu_str = [str(i) for i in yewu_and_jishu]
    df = pd.DataFrame.from_records(levels, columns=[i[0] for i in patterns])
    df['value'] = yewu_and_jishu_str

    return df

def yewu_yu_jishu_chapter(filename, total_doc_content, sub_dir):
    start, end = None, None
    for idx, line in enumerate(total_doc_content):
        if not isinstance(line, str):
            continue
        line = line.strip()
        if re.search(r'^第.{0,2}[节章]\s*业务.{0,2}技术',  line):
            start = idx
            continue
        if re.search(r'^第.{0,2}[节章]', line) and start is not None:
            end = idx
            break
    if start is not None and end is not None:
        yewu_and_jishu = total_doc_content[start:end]
    else:
        print(f'第六节 业务与技术 没找到 in {filename}')
        yewu_and_jishu = []

    df = get_hierarchy_of_chapter(yewu_and_jishu)
    save_path = os.path.join('最新文件抽取结果', '第六节业务与技术', sub_dir, filename + '.xlsx')
    os.makedirs(os.path.dirname(save_path), exist_ok=True)
    df.to_excel(save_path)
    return yewu_and_jishu


########################################################################################################################
def process_all_file_concurrent(filename, file_batch, root, sub_dir):
    file_batch = [os.path.join(root, sub_dir, i) for i in file_batch]
    total_doc_content = read_merged_doc(file_batch)
    # # 业务与技术章节提取
    # result_yewu_and_jishu = yewu_yu_jishu_chapter(filename, total_doc_content, sub_dir)
    #经营成果分析
    result_jingyingchengguo = extract_table_given_key_yingyeshouwu_analize(filename, total_doc_content)
    # 主要财务数据和财务指标
    result_zhuyaocaiwushuju = extract_table_given_key_zhuyaocaiwushuju(filename, total_doc_content)
    # 主营业务构成
    result_zhuyingyewugoucheng = extract_table_given_key_zhuyingshourugoucheng(filename, total_doc_content)
    # 本次发行概况
    result_bencifaxinggaikuang = extract_table_given_key_bencifaxinggaikuang(filename, total_doc_content)
    # 发行人及本次发行的中介机构基本情况
    result_faxingrenandfaxingzhongjie = extract_table_given_key_faxingrenhefaxingzhongjieV2(filename, total_doc_content)
    # 财务会计信息与管理层分析章节抽取
    result_caiwukuaiji = extract_table_given_key_caiwukuaiji_xinxi(filename, total_doc_content)
    # 发行前后公司股本变化
    result_guben_bianhua = extract_table_given_key_guben_bianhua(filename, total_doc_content)
    # 员工专业结构、学历结构
    result_yuangong_jiegou = extract_table_given_key_yuangong_zhuanye_xueli(filename, total_doc_content)
    # 报告期内前五大客户销售情况
    result_top5_kehu_xiaoshou = extract_table_given_key_top_five_kehu(filename, total_doc_content)
    # 报告期内前五大供应商
    result_top5_gongyingshang = extract_table_given_key_top_five_gongyings(filename, total_doc_content)
    # 主营业务收入，按产品、地域等划分
    result_zhuyingshouru_different_part = extract_table_given_key_zhuying_different_part(filename, total_doc_content)
    # 毛利和毛利率分析
    result_maolilv = extract_table_given_key_maolilv(filename, total_doc_content)
    # 前十大股东持股情况
    result_top10_gudong = extract_table_given_key_top10_gudong(filename, total_doc_content)
    # 发行人符合科创属性、行业领域要求
    result_meet_demand = extract_table_given_key_meet_demand(filename, total_doc_content)
    result = {
        '主要财务数据和财务指标':result_zhuyaocaiwushuju,
        '主营业务构成':result_zhuyingyewugoucheng,
        '本次发行概况':result_bencifaxinggaikuang,
        '发行人及本次发行的中介机构基本情况':result_faxingrenandfaxingzhongjie,
        '财务会计信息与管理层分析章节抽取':result_caiwukuaiji,
        '发行前后股本变化情况':result_guben_bianhua,
        '员工专业学历':result_yuangong_jiegou,
        '前五大客户销售':result_top5_kehu_xiaoshou,
        '前五大供应商':result_top5_gongyingshang,
        '经营成果分析': result_jingyingchengguo,
        "毛利率":result_maolilv,
        "前十大股东":result_top10_gudong,
        "符合科创属性及行业领域要求":result_meet_demand,
        '主营业务收入，按产品、地域等划分':result_zhuyingshouru_different_part,
        # '业务与技术':result_yewu_and_jishu
    }

    write_to_excel(result, filename, sub_dir)
    return

def main():
    multi_p = 10
    # multi_p = False

    root = r'C:\Users\hujunchao\Documents\PdfDir\original\招股书docx'
    sub_dirs = [
        'batch1-主',
        'batch2-主',
        'batch3-主',
        'batch4-主',
        'batch5-主',
        'batch6-主',
        'batch7-主',
        'batch8-主',
    ]
    # sub_dirs = ['batch7-主', 'batch7-次']
    # sub_dirs = ['batch1', 'batch1-次']
    for sub_dir in sub_dirs:
        all_files = os.listdir(os.path.join(root, sub_dir))
        all_files = [i for i in all_files if not re.match(r'~$', i)]
        # all_files = [i for i in all_files if i.startswith('zg-688004')]
        merged = merge_files(all_files)
        merged = [(filename, file_batch, root, sub_dir) for filename, file_batch in merged
                  # if filename == '[688506] 百利天恒首次公开发行股票并在科创板上市招股说明书'
                  ]
        if multi_p:
            pool = Pool(multi_p)
            pool.starmap(process_all_file_concurrent, merged)
        else:
            [process_all_file_concurrent(*i) for i in merged]


if __name__ == '__main__':

    main()